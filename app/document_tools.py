from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

def extract_document(path: Path):
    ext=path.suffix.lower()
    if ext=='.pdf':
        reader=PdfReader(str(path)); return [(i+1, (page.extract_text() or '').strip()) for i,page in enumerate(reader.pages)]
    if ext=='.docx':
        doc=Document(str(path)); text=[]
        for para in doc.paragraphs:
            t=para.text.strip()
            if t: text.append(t)
        for table in doc.tables:
            for row in table.rows:
                vals=[c.text.strip() for c in row.cells]
                if any(vals): text.append(' | '.join(vals))
        return [(1,'\n'.join(text))] if text else [(1,'')]
    if ext=='.pptx':
        prs=Presentation(str(path)); pages=[]
        for i,slide in enumerate(prs.slides,1):
            bits=[]
            for shape in slide.shapes:
                if hasattr(shape,'text') and shape.text.strip(): bits.append(shape.text.strip())
                if getattr(shape,'has_table',False):
                    for row in shape.table.rows:
                        bits.append(' | '.join(cell.text.strip() for cell in row.cells))
            pages.append((i,'\n'.join(bits)))
        return pages or [(1,'')]
    raise ValueError('Unsupported file type')
